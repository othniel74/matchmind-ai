from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/product'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.04
styles['Title'].font.size = Pt(28)
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Subtitle'].font.size = Pt(12)
styles['Heading 1'].font.size = Pt(20)
styles['Heading 1'].paragraph_format.space_after = Pt(12)
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].paragraph_format.space_before = Pt(13)
styles['Heading 2'].paragraph_format.space_after = Pt(6)
header = sec.header.paragraphs[0]
header.text = 'PITCHLIGHT   /   PRODUCT STRATEGY'
header.style = styles['Normal']
header.runs[0].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.text = '18 October 2026   /   Working strategy                                      '
footer.runs[0].font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text, bold=False):
    x = doc.add_paragraph()
    x.add_run(text).bold = bold
    return x

def h(text):
    doc.add_heading(text, 2)

def page(title):
    title_para = doc.add_heading(title, 1)
    title_para.paragraph_format.page_break_before = True

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, width in zip(t.columns, widths): cell.width = Inches(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = text
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '243C49' if ri == 0 else ('F0F4F6' if ri % 2 else 'FFFFFF')); props.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e = OxmlElement('w:'+edge); e.set(qn('w:w'), '95'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            props.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(10)
                    if ri == 0: run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(3)
    gap.paragraph_format.line_spacing = Pt(4)
    gap.add_run('').font.size = Pt(4)

doc.add_heading('Pitchlight product strategy', 0)
doc.add_paragraph('Production vision and delivery priorities', style='Subtitle')
p('Prepared for the Pitchlight product and engineering team. This document consolidates our decisions about product value, real datasets, graph intelligence, production readiness and the hackathon milestone. Status reflects the local implementation discussed on 18 October 2026.')
h('The product direction')
p('We are building Pitchlight for sustained production use. A successful hackathon demonstration is an early milestone; adoption depends on helping users make better editorial decisions with less effort and with evidence they can inspect.')
p('Our proposed product is a visual match-intelligence and storytelling workspace. It should identify a meaningful moment, explain why it matters, show the supporting evidence and help a producer turn it into an audience-appropriate story.')
p('This is a product hypothesis. We have not established competitive uniqueness, customer demand or production readiness. Those claims need comparison with alternatives and observation of prospective users doing real work.')
h('Who should use it first')
table(['User', 'Recurring task', 'Expected benefit'], [
 ('Broadcast producers and match analysts', 'Find, verify and prepare a useful match story', 'Less searching and a shorter path to an approved output'),
 ('Coaches and club analysts', 'Review sequences, team shape and passing relationships', 'Visual evidence for performance discussions'),
 ('Players and academies', 'Understand positioning and decisions', 'Clearer review of particular actions'),
 ('Fans and educators', 'Understand how and why a moment developed', 'Accessible explanations tailored to interests')], [1.65,2.45,2.8])
p('Recommended initial adoption audience: broadcast producers and match analysts. This matches the Studio interface and gives us a specific workflow to validate. The other audiences are expansion opportunities, not four simultaneous product commitments.', True)

page('A distinctive editorial workflow')
p('The proposed differentiator is the connection between discovery, evidence and publication. Charts and fluent text alone do not establish a useful product. The system must help a user decide what to say, understand the support for it and control what reaches an audience.')
h('An example worth testing')
p('A team has been progressing through a particular midfielder. During a later passage, that player receives fewer completed passes and the usual passing route becomes less active. Pitchlight identifies the change and retrieves the relevant sequences. If synchronized tracking is available, it examines the nearby passing options and spacing.')
p('The system must distinguish observation from explanation. A decline in passing connections does not by itself prove that an opponent caused it. A story should state the measured change, describe the available spatial evidence and qualify any suggested cause.')
h('From event to approved story')
for text in [
 '1  Detect a candidate moment using source events, graph measurements and available tracking.',
 '2  Assemble an evidence packet with source IDs, timestamps, comparison windows, data quality and calculation details.',
 '3  Draft a Why this moment matters story card with the observation, relevance, supporting replay and limitations.',
 '4  Adapt the wording for a broadcaster, analyst or fan and the selected language.',
 '5  Let a producer inspect, edit, approve or reject the draft before scheduling an overlay or recap.',
 '6  Track publication versions and flag affected stories if the source is corrected.'
]: p(text)
h('The first signature feature')
p('Build an evidence-backed Why this moment matters card synchronized with the selected match. It should offer a concise explanation, a selectable replay moment, source evidence, a quality indicator and an editable producer draft. Editorial approval and automatic correction propagation are proposed capabilities, not verified features of the current app.')
h('What a useful story must contain')
p('A specific claim; an appropriate comparison or match context; inspectable evidence; a clear distinction between measured facts and interpretation; an audience-specific output; and a producer decision. Unsupported observations should be withheld or labeled as uncertain, rather than dressed up as tactical conclusions.')

page('What the datasets contribute')
p('Real samples help us test movement and analytical calculations against recorded football situations. They do not automatically make a finding correct, confer footage rights or validate the product as a whole. Matches from different providers must remain separate unless a reliable identity and timing alignment is established.')
table(['Source', 'Added locally', 'What it enables'], [
 ('Synthetic scenario', 'Deterministic match events, scoring, recipients and streamed replay', 'Reproducible tests and the core hackathon demonstration'),
 ('SkillCorner', 'A two-minute tracking excerpt from Melbourne Victory v Auckland, 17 May 2025', 'Recorded player and ball movement, spacing and team width and depth'),
 ('Metrica Sports', '3,001 frames at 25 Hz and 35 synchronized events from anonymous Sample Game 1', 'Event-to-replay inspection, source-based counts and a goal update at the shot end frame'),
 ('StatsBomb', '4,407 events from Argentina v France, World Cup final 2022', 'Passing networks, shot locations, source xG and event evidence')], [1.15,2.85,2.9])
h('The important boundary today')
p('The real-data views use their own matches and clocks. The surrounding synthetic dashboard, narrative queue and overlays are not a unified analysis of the selected research match. The app labels this separation. A production version needs one selected match context to drive every panel and published output.')
h('Movement and visual realism')
p('Tracking supplies recorded coordinates, including provider estimates where available. It does not provide exact player likenesses, body motions or licensed match video. The current character models, kits and running animations are illustrative. SkillCorner distinguishes detected and extrapolated positions; Metrica lacks those quality flags and ball height. Missing positions should remain unavailable rather than be invented.')
h('Rights and commercial access')
p('The SkillCorner MIT notice is retained. Metrica attribution and its source usage notice are retained without assuming blanket commercial permission. StatsBomb is configured for local noncommercial research under its published agreement. The new StatsBomb and Metrica source files are excluded from Git and Docker deployment assets. Production use needs provider-specific permissions for access, display, derived outputs, retention and redistribution. Live feeds and match footage require their own arrangements.')

page('NetworkX and the knowledge graph')
h('NetworkX discovers measurable patterns')
p('A passing graph represents players as nodes and passes as directed connections. NetworkX can measure connection strength, passing involvement and changes between time windows. This helps identify dominant partnerships, dependence on particular players and routes that become less active. Advanced measures must be evaluated for stability and football relevance before they become product claims.')
p('Our StatsBomb implementation uses completed passes with explicit recipients. Each weighted connection retains its source event IDs. Displayed node locations are mean event locations, not tracking positions or a formation. The synthetic analysis also uses recipient-based passing networks, with a labeled proxy fallback for older fixtures.')
h('The knowledge graph preserves relationships and context')
p('A production knowledge graph would link matches, periods, teams, players, possessions, events, observations, evidence packets, story versions and publication decisions. For example: a player made a pass to a teammate during a possession that led to a shot; an observation cites that sequence; an approved story explains the observation.')
p('It should answer which sequences support a claim, what changed after a substitution and which outputs are affected by a corrected event. A graph database is justified when these recurring relationship queries improve the workflow. It is not a substitute for validated measurements or a requirement merely to enlarge the technology stack.')
table(['Responsibility', 'NetworkX', 'Knowledge graph'], [
 ('Primary role', 'Compute network measurements', 'Store and retrieve relationships and provenance'),
 ('Typical input', 'Validated passes and time windows', 'Events, identities, observations and story versions'),
 ('Typical output', 'A measured change or candidate pattern', 'Supporting sequences and relevant context'),
 ('Correction handling', 'Recompute affected measurements', 'Find and invalidate affected observations and stories')], [1.45,2.7,2.75])
h('How they work together')
p('A graph measurement detects a change. Tracking, when available, adds spatial evidence. The knowledge graph retrieves the relevant events and context. AI drafts a qualified explanation from that evidence packet. A producer reviews it and controls publication. This complete flow remains a production target; it is not already verified end to end.')
p('NetworkX analysis works locally. Existing graph database adapters have limited mocked verification; a production knowledge graph, live database behavior and correction propagation still require implementation and testing.')

page('Production foundations and current gaps')
p('The present system is a working local prototype. Production readiness means repeatable behavior, controlled access, durable state and trustworthy outputs under incomplete, delayed and corrected data. A polished interface and passing unit tests are necessary evidence, but are not sufficient proof of readiness.')
table(['Area', 'Verified foundation', 'Work still required'], [
 ('Data and time', 'Synthetic stream, reconnect resume and separate research adapters', 'One match context, period-aware clocks, stable IDs, deduplication, late events and corrections'),
 ('Persistence', 'SQLite replay persistence and integrity checks', 'Production database design, migrations, backups, recovery and concurrency'),
 ('AI and language', 'Evidence-checked deterministic narratives; English and Spanish modes', 'Verified Foundry inference, orchestration traces, multilingual evaluation and full recaps'),
 ('Editorial workflow', 'Studio preview, timed overlays and presentation controls', 'Editable drafts, approval permissions, scheduling, version history and withdrawal'),
 ('Identity and access', 'Local workspace exit control', 'Authentication, workspaces, roles, tenant isolation and audit trails'),
 ('Operations', 'Local builds and automated checks', 'Deployment verification, monitoring, incident handling, load tests and cost controls')], [1.05,2.65,3.2])
h('One match should drive every output')
p('Introduce a canonical match context with provider identity, period and clock, event version, provenance and capability flags. Statistics, movement, evidence, stories, overlays and recaps should all subscribe to this context. Event-only sources should not pretend to have continuous tracking, and unavailable metrics should remain explicitly unavailable.')
h('AI needs a controlled publication boundary')
p('Compute factual quantities in deterministic code and supply evidence to the model. Validate generated claims, numbers and source IDs before publication. Keep the measured facts separate from interpretation, record model and prompt versions and test failure cases. Claims of confidence need calibration; a low data-quality flag is not a calibrated probability.')
p('Live Foundry calls and Azure hosting are unverified. Local SDK construction and deterministic fallback are not evidence of live AI. Today’s scope remains local; Azure setup and commercial provider agreements are later activities, not implicit authorization to deploy or spend.')

page('Delivery roadmap and adoption evidence')
p('Prioritize a complete workflow that users can evaluate. The sequence below is proposed; it is not a committed schedule or a promise that production readiness can be completed within the hackathon.')
table(['Stage', 'Deliverable', 'Evidence required to move on'], [
 ('1  Shared match context', 'One source and clock across analysis, replay and story evidence', 'No cross-match leakage; correct seeking and source correction behavior'),
 ('2  Signature story card', 'Why this moment matters with replay evidence and editable wording', 'Users can inspect the support and make an editorial decision'),
 ('3  Recaps and personalization', 'Halftime and full-time recaps; audience, language and player interests', 'Grounded outputs without future-event leakage; distinct user experiences'),
 ('4  Verified AI workflow', 'Foundry execution with useful specialized roles and publication gates', 'Real traces, grounded outputs, failure recovery and evaluated latency'),
 ('5  Production pilot', 'Authorized feed, authenticated workspace and durable editorial state', 'Recovery, isolation, monitoring, rights and support processes verified'),
 ('6  Submission and rollout', 'Accessible demo, documentation and staged release', 'Reproducible setup and honest, tested product claims')], [1.35,2.7,2.85])
h('How we test whether users want it')
p('Observe broadcast producers and match analysts using the prototype on representative tasks. Ask them to find a meaningful sequence, inspect its evidence, edit the explanation and prepare an output. Compare this with their current workflow before setting improvement targets.')
table(['Question', 'Measure'], [
 ('Are suggestions useful?', 'Accepted, edited and rejected stories with reasons'),
 ('Does it save effort?', 'Time from event discovery to approved story against a baseline'),
 ('Can users trust it?', 'Factual error rate, unsupported claims and correction handling'),
 ('Will they return?', 'Repeated use, completed review tasks and willingness to pilot')], [2.15,4.75])
p('No adoption, accuracy, time-saving or uniqueness targets have yet been validated. Set thresholds with pilot users, include unsuccessful cases and avoid treating a compelling demonstration as proof of demand.')

page('The hackathon milestone and decisions')
p('Pitchlight aligns with the Inside the Game challenge as a match-intelligence and storytelling studio. The hackathon should demonstrate a focused production-shaped workflow on synthetic football-realistic events. Real-data research remains supporting work, subject to source permissions.')
h('What the submission still needs')
p('Complete the synthetic event-to-explanation-to-overlay flow and automated recaps. Demonstrate genuinely different audience experiences. Verify Microsoft AI execution and useful agent handoffs before describing them as working. Provide a reproducible public code repository, a working demo or test build, testing instructions and an English pitch with a demonstration video under two minutes. Cloud hosting is one access option; the rules also permit a functioning demo or test build.')
p('The official rules specify synthetic data, third-party permissions and five equally weighted judging areas: technological implementation, agentic design and innovation, real-world impact, user experience and presentation, and category adherence. Listing Microsoft technologies is not equivalent to demonstrating their value. We should select a category only when the implementation supports it.')
h('Dates to manage')
table(['Milestone', 'Official Pacific time'], [
 ('Registration closes', '20 October 2026 at 12 noon'),
 ('Project submission closes', '27 October 2026 at 11 59 pm'),
 ('Judging period ends', '10 November 2026 at 11 59 pm')], [3.25,3.65])
h('Decisions already established')
p('Build for production adoption beyond the hackathon. Keep current work local and use free assets and samples. Make synthetic data the submission’s core demonstration. Keep research matches separate. Prioritize a unified match context and evidence-backed storytelling over further cosmetic realism. Treat producer control, source provenance and commercial data permissions as product requirements.')
h('Open decisions')
p('Validate the initial customer segment and editorial workflow; define the paid-product value proposition; choose an authorized production data provider; establish pilot success thresholds; choose a category and Microsoft architecture; and agree the security, reliability and support expectations for a first production pilot.')
page('Sources and verification record')
p('The discussion and local implementation records are the basis for this working strategy. They describe a point in time and do not certify production readiness, commercial rights or eligibility. Recheck the official rules and provider terms before submission or distribution.')
h('References and status evidence')
for text in [
 'Official rules: https://github.com/microsoft/insidethegamehackathon/blob/main/OFFICIAL%20RULES.md',
 'Data sources: github.com/SkillCorner/opendata; github.com/metrica-sports/sample-data; github.com/statsbomb/open-data. Retained provider terms govern use.',
 'Local implementation evidence: docs/REAL_DATA_SOURCES.md; docs/REAL_TRACKING_REPLAY.md; docs/VERIFICATION_2026_10_17.md; docs/RELEASE_CHECKLIST.md. Older milestone notes are historical; this document does not replace executable verification.'
]:
    x = p(text)
    for run in x.runs: run.font.size = Pt(9)

h('What the current checks establish')
table(['Evidence', 'Scope and limit'], [
 ('Backend verification', '55 backend tests passed across the suite and a rerun of the storage test after resolving a temporary-directory setup issue. This does not establish live provider or cloud behavior.'),
 ('Tracking checks', 'Coordinate mapping and interpolation were checked across 1,201 SkillCorner frames and 3,001 Metrica frames. Goal timing and ball-display height calculations have targeted checks.'),
 ('Browser and build checks', 'Local research controls, source switching, evidence selection and production frontend builds were verified. This is not a production load or accessibility certification.'),
 ('AI and graph database status', 'Optional SDK construction and mocked graph adapter tests exist. Live Foundry inference, real database behavior and end-to-end correction propagation remain unverified.')], [1.65,5.25])
h('How to use this document')
p('Use the product hypothesis to choose the next vertical slice and the roadmap to define its acceptance evidence. Keep status claims tied to dated tests. Update this document after customer interviews, provider agreements, live AI verification and a production pilot; do not carry prototype assumptions into a commercial launch unchanged.')

doc.core_properties.title = 'Pitchlight product strategy'
doc.core_properties.subject = 'Production vision adoption graph intelligence and delivery roadmap'
doc.core_properties.author = 'Pitchlight'
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.save(OUT / 'Pitchlight_AI_Product_Strategy.docx')
print(OUT / 'Pitchlight_AI_Product_Strategy.docx')
