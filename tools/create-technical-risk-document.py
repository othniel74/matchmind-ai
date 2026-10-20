from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ROW_HEIGHT_RULE

ROOT=Path(__file__).resolve().parents[1]
source=ROOT/'docs/TECHNICAL_IMPLEMENTATION_AND_RISK_CONTROLS.md'
out=ROOT/'docs/technical/MatchMind_AI_Technical_Implementation_and_Risk_Controls.docx'
out.parent.mkdir(parents=True,exist_ok=True)
doc=Document(ROOT/'docs/product/MatchMind_AI_Product_Strategy_v2.docx')
for element in list(doc.element.body):
    if element.tag!=qn('w:sectPr'):doc.element.body.remove(element)
for section in doc.sections:
    section.header.paragraphs[0].text='MATCHMIND AI   /   TECHNICAL IMPLEMENTATION'
    section.header.paragraphs[0].runs[0].font.size=Pt(9)
    for p in section.footer.paragraphs:
        for r in p.runs:
            if 'Working strategy' in r.text:r.text=r.text.replace('Working strategy','Technical specification')
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.04
normal.paragraph_format.space_after=Pt(6)

def table(rows):
    headers=rows[0]; body=rows[2:]
    widths=[1.05,2.8,3.05] if headers[0].startswith('ID') else [1.4,2.8,2.7]
    if len(headers)==2:widths=[1.5,5.4]
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    t.rows[0].height=Inches(.32)
    t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    for values in body:
        for c,v in zip(t.add_row().cells,values):c.text=v
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci])
            props=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'243C49' if ri==0 else ('F0F4F6' if ri%2 else 'FFFFFF'));props.append(shading)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.02
                for r in p.runs:
                    r.font.size=Pt(10)
                    if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(4);p.paragraph_format.space_after=Pt(3)

lines=source.read_text(encoding='utf-8').splitlines()
first=True;new_page=False;i=0;list_num=None
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- page -->':new_page=True;i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_heading(line[2:],0 if first else 1)
        if new_page:p.paragraph_format.page_break_before=True
        first=False;new_page=False
    elif line.startswith('## '):doc.add_heading(line[3:],2)
    elif len(line)>2 and line[0].isdigit() and line[1:3]=='. ':
        if line.startswith('1. '):
            numbering=doc.part.numbering_part.element
            style_num=doc.styles['List Number'].element.pPr.numPr.numId.val
            base=numbering.num_having_numId(style_num)
            num=numbering.add_num(base.abstractNumId.val)
            num.add_lvlOverride(ilvl=0).add_startOverride(1)
            list_num=num.numId
        paragraph=doc.add_paragraph(line[3:],style='List Number')
        if list_num is not None:
            props=paragraph._p.get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val=0
            props.get_or_add_numId().val=list_num
    else:
        p=doc.add_paragraph(line)
    i+=1
for style in ['Title','Subtitle','Heading 1','Heading 2']:
    doc.styles[style].font.color.rgb=RGBColor(0,0,0)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='MatchMind AI technical implementation and risk controls'
doc.core_properties.subject='Production architecture risk register acceptance tests and monitoring'
doc.core_properties.comments='Consolidated specification. Implemented partial planned and unverified controls are distinguished. Earlier documents preserved.'
doc.save(out)
print(out)
