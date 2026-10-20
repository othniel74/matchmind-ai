from pathlib import Path
from docx import Document
from docx.shared import Pt

root = Path(__file__).resolve().parents[1]
source = root/'docs/product/MatchMind_AI_Product_Strategy.docx'
target = root/'docs/product/MatchMind_AI_Product_Strategy_v2.docx'
doc = Document(source)

replacements = {
 'It should answer which sequences support a claim':
 'The knowledge graph is a core production component, designed alongside the shared match pipeline. Its schema, stable identities, evidence links and event versions belong in the first stage. Database selection and deployment can progress in stages; the evidence relationships must exist before AI explanations and publication depend on them.',
 'A graph measurement detects a change.':
 'Validated match events feed NetworkX analysis and a knowledge graph of context and evidence. AI retrieves the relevant measurements and sequences, drafts an explanation and passes it to producer review. Approved versions drive overlays and recaps. NetworkX computes patterns; the knowledge graph preserves their support and downstream dependencies.',
 'Introduce a canonical match context':
 'Introduce a canonical match context and knowledge graph foundation together: provider identity, period and clock, stable event IDs, source versions and provenance. Connect observations, evidence, story versions and published outputs to that context. Every panel must use the selected match; event-only sources must keep tracking-dependent metrics unavailable.',
 'Build for production adoption beyond the hackathon. Keep current work local':
 'Build for production adoption beyond the hackathon. Keep current work local and use free assets and samples. Make synthetic data the submission core and keep research matches separate. Treat the knowledge graph and shared match context as core foundations, with evidence-backed storytelling, producer control, correction handling and commercial data permissions.'
}
for para in doc.paragraphs:
    for prefix, text in replacements.items():
        if para.text.startswith(prefix):
            para.text = text
            break

for table in doc.tables:
    heads = [c.text for c in table.rows[0].cells]
    if heads[:2] == ['Stage','Deliverable']:
        rows = [
          ('1  Match and graph foundations','Shared match context, graph schema, stable IDs and evidence links','No cross-match leakage; versioned sources and retrievable evidence'),
          ('2  Signature story card','Graph-backed evidence retrieval, replay and editable explanation','Every claim links to evidence; a producer can review the draft'),
          ('3  Stories and personalization','Developing-story timeline, recaps and audience or player interests','Supported updates, no future leakage and distinct user experiences'),
          ('4  Verified AI and graph queries','Foundry execution, Ask the match and tested graph storage','Real traces, grounded answers, useful handoffs and query tests'),
          ('5  Production pilot','Authorized feed, secure workspace and graph correction propagation','Affected stories flagged; recovery, isolation and monitoring verified'),
          ('6  Submission and rollout','Accessible demo, documented architecture and staged release','Reproducible setup, permissions review and honest product claims')]
        for row, values in zip(table.rows[1:],rows):
            for cell,text in zip(row.cells,values):
                cell.paragraphs[0].runs[0].text=text
                for run in list(cell.paragraphs[0].runs)[1:]:run.text=''
    if heads[:2] == ['Area','Verified foundation']:
        for row in table.rows[1:]:
            if row.cells[0].text == 'Persistence':
                row.cells[2].paragraphs[0].runs[0].text='Production storage, knowledge graph schema, migrations, backups, recovery and concurrency'

heading=doc.add_heading('Knowledge graph delivery plan',1)
heading.paragraph_format.page_break_before=True

def h(text):doc.add_heading(text,2)
def p(text):doc.add_paragraph(text)

p('Revision 2 makes the knowledge graph explicit in the core architecture and roadmap. It also places Ask the match, developing stories and correction handling around the same evidence model. These are planned production capabilities; the current graph adapters do not establish their completion.')
h('Core relationships')
p('Model Match, Period, Team, Player, Possession, EventVersion, Observation, EvidencePacket, StoryVersion, Approval and PublishedOutput. Events belong to a match and period; observations cite event versions; story versions explain observations; approvals authorize specific versions; published outputs reference the approved story and scheduling context.')
p('Keep raw provider records in durable source storage. Use the graph to preserve relationships and provenance, and NetworkX to compute network measurements from validated projections. Record the algorithm version, comparison window and inputs behind each observation.')
h('Features built on the graph')
p('Why this moment matters retrieves the supporting sequences and measurements for an editable story. Ask the match answers from authorized evidence in the selected match and clock window. The developing-story timeline records when an observation strengthens, weakens or becomes outdated. Recaps reuse verified moments while retaining evidence links.')
h('Corrections and editorial control')
p('A corrected source creates a new event version. Identify affected measurements and recompute them, then flag dependent observations, story drafts, scheduled overlays and recaps for review. Preserve the prior version and producer decisions. Published material needs a recorded correction or withdrawal decision; it must not silently change.')
h('Acceptance evidence')
for text in [
 'Each displayed claim resolves to the correct match, event versions, calculation and replay interval.',
 'Queries respect the selected match clock and workspace permissions, without future-event leakage.',
 'A corrected pass identifies every dependent observation and story and triggers the defined review state.',
 'Graph queries and failure recovery are verified against a real database; mocked adapter tests are insufficient.',
 'Database choice is based on demonstrated query, operational and cost needs. The graph remains in the design while that choice is validated.'
]:
    para=doc.add_paragraph(text,style='List Bullet')
    para.paragraph_format.space_after=Pt(5)

doc.core_properties.comments='Revision 2 explicitly includes the knowledge graph in the core architecture and delivery roadmap. Original version retained.'
doc.save(target)
print(target)
